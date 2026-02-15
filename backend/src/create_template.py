from docx import Document
import os
from pathlib import Path

# Define path
BASE_DIR = Path(__file__).resolve().parent.parent
TEMPLATE_DIR = BASE_DIR / 'templates' / 'reports'
TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_FILE = TEMPLATE_DIR / 'informe_propiedad_fixed.docx'

def create_template():
    doc = Document()
    
    doc.add_heading('Informe de Propiedad', 0)
    
    doc.add_paragraph('Generado para: {{ nombre_cliente }}')
    
    doc.add_heading('Detalles de la Propiedad', level=1)
    
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    
    # Header
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Campo'
    hdr_cells[1].text = 'Valor'
    
    # Data rows
    data = [
        ('Dirección', '{{ address }}'),
        ('Precio', '{{ price }}'),
        ('Superficie Total', '{{ surface_total }}'),
        ('Superficie Útil', '{{ surface_useful }}'),
        ('Dormitorios', '{{ bedrooms }}'),
        ('Baños', '{{ bathrooms }}'),
        ('Estacionamientos', '{{ parking_spots }}'),
        ('Bodegas', '{{ storage_units }}'),
    ]
    
    for label, tag in data:
        row_cells = table.add_row().cells
        row_cells[0].text = label
        row_cells[1].text = tag
        
    doc.add_paragraph('\nEste documento fue generado automáticamente.')
    
    doc.save(OUTPUT_FILE)
    print(f"Template created successfully at: {OUTPUT_FILE}")

if __name__ == "__main__":
    create_template()
